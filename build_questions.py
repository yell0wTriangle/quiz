"""Convert the complete quiz worksheet in the supplied workbook into browser data."""
import json
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path("IT Questions (2).xlsx")
SHEET = "IT Questions for Quiz"
TARGET = Path("questions.js")


def text(value):
    return "" if value is None else str(value).strip()


workbook = load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook[SHEET]
questions = []

for row in sheet.iter_rows(min_row=2, values_only=True):
    number = row[0]
    if not isinstance(number, (int, float)):
        continue  # Section labels and blank rows.
    question = {
        "number": int(number),
        "topic": text(row[1]),
        "question": text(row[2]),
        "options": [text(value) for value in row[3:7]],
        "answer": text(row[7]).upper()[:1],
    }
    while question["options"] and not question["options"][-1]:
        question["options"].pop()
    questions.append(question)

expected = list(range(1, 669))
actual = [item["number"] for item in questions]
if actual != expected:
    raise SystemExit(f"Could not parse all questions. Parsed {len(actual)} rows; expected 668.")
if any(not item["topic"] or not item["question"] or len(item["options"]) < 2 or item["answer"] not in "ABCD" for item in questions):
    raise SystemExit("The workbook contains incomplete quiz rows.")

TARGET.write_text(
    "// Generated from IT Questions (2).xlsx. Run build_questions.py after updating the workbook.\n"
    + "window.QUIZ_QUESTIONS = "
    + json.dumps(questions, ensure_ascii=False, separators=(",", ":"))
    + ";\n",
    encoding="utf-8",
)
print(f"Wrote {len(questions)} questions to {TARGET}")
